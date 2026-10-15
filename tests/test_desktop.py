import asyncio
from datetime import datetime, timezone
import tempfile
import time
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch
from zipfile import ZipFile

from product_council.desktop_session import (
    DesktopProvider, ModelChoice, ModelDiscoveryError, discover_models, run_session, safe_error, validate_choice,
)
from product_council.providers import LLMRequest, MockProvider
from product_council.word_report import save_word
from tests.test_council import QUESTION, _brief, _config, _run
from product_council import Council
from product_council.connections import ConnectionStore, ConnectionStoreError


class MemoryVault:
    def __init__(self):
        self.values = {}

    def get_password(self, service, username):
        return self.values.get((service, username))

    def set_password(self, service, username, password):
        self.values[service, username] = password

    def delete_password(self, service, username):
        del self.values[service, username]


class ConnectionStoreTests(unittest.TestCase):
    def test_save_reload_switch_and_forget_independent_providers(self):
        vault = MemoryVault()
        store = ConnectionStore(vault)
        openai = ModelChoice("openai", "gpt-model", "openai-test-key")
        anthropic = ModelChoice("anthropic", "claude-model", "anthropic-test-key")
        self.assertIsNone(store.load())
        store.save(openai)
        store.save(anthropic)
        restored = ConnectionStore(vault)
        self.assertEqual(restored.load(), anthropic)
        self.assertEqual(restored.load("openai"), openai)
        restored.forget("anthropic")
        self.assertEqual(restored.load(), openai)
        self.assertIsNone(restored.load("anthropic"))
        restored.forget("openai")
        self.assertIsNone(restored.load())
        self.assertEqual(vault.values, {})

    def test_storage_errors_are_explicit_and_never_include_keys(self):
        from unittest.mock import Mock

        vault = Mock()
        vault.get_password.side_effect = RuntimeError("private-key-in-error")
        store = ConnectionStore(vault)
        for operation in (lambda: store.load(), lambda: store.save(ModelChoice("openai", "m", "key")),
                          lambda: store.forget("openai")):
            with self.assertRaises(ConnectionStoreError) as caught:
                operation()
            self.assertNotIn("private-key-in-error", str(caught.exception))
        vault.set_password.assert_not_called()

    def test_invalid_stored_data_is_not_used_or_overwritten(self):
        vault = MemoryVault()
        vault.set_password(ConnectionStore.SERVICE, ConnectionStore.ACCOUNT, '{"unexpected": "data"}')
        store = ConnectionStore(vault)
        with self.assertRaises(ConnectionStoreError):
            store.load()
        with self.assertRaises(ConnectionStoreError):
            store.save(ModelChoice("openai", "m", "key"))
        self.assertEqual(vault.get_password(store.SERVICE, store.ACCOUNT), '{"unexpected": "data"}')


class DesktopSessionTests(unittest.TestCase):
    def choice(self):
        return ModelChoice("openai", "customer-chosen-model", "test-key-not-a-real-secret")

    def test_model_choice_validates_required_values_and_azure_endpoint(self):
        for choice in (ModelChoice("mock", "m", "k"), ModelChoice("openai", "", "k"),
                       ModelChoice("openai", "m", " "), ModelChoice("azure_openai", "m", "k"),
                       ModelChoice("azure_openai", "m", "k", "http://example.com"),
                       ModelChoice("azure_openai", "m", "k", "https://user:pass@example.com")):
            with self.subTest(provider=choice.provider), self.assertRaises(ValueError):
                choice.spec()
        azure = ModelChoice("azure_openai", "deployment", "k", "https://resource.openai.azure.com")
        self.assertFalse(azure.spec()["use_entra_id"])
        self.assertNotIn(self.choice().api_key, repr(self.choice()))

    def test_validation_checks_real_completion_and_closes_provider(self):
        provider = AsyncMock()
        provider.complete.return_value = '{"ok": true}'
        with patch("product_council.desktop_session.create_provider", return_value=provider) as create:
            asyncio.run(validate_choice(self.choice()))
        self.assertEqual(create.call_args.args[0]["api_key"], self.choice().api_key)
        request = provider.complete.call_args.args[0]
        self.assertEqual(request.model, self.choice().model)
        self.assertNotIn(self.choice().api_key, request.user)
        provider.close.assert_awaited_once()

    def test_discovery_does_not_require_model_or_make_inference_request(self):
        choice = ModelChoice("openai", "", "key")
        provider = AsyncMock()
        provider.list_models.return_value = ["gpt-example"]
        with patch("product_council.desktop_session.create_provider", return_value=provider):
            self.assertEqual(asyncio.run(discover_models(choice)), ["gpt-example"])
        provider.complete.assert_not_awaited()
        provider.close.assert_awaited_once()

    def test_discovery_reports_empty_catalog_and_closes_provider(self):
        provider = AsyncMock()
        provider.list_models.return_value = []
        with patch("product_council.desktop_session.create_provider", return_value=provider), \
                self.assertRaisesRegex(ModelDiscoveryError, "no supported chat models"):
            asyncio.run(discover_models(self.choice()))
        provider.close.assert_awaited_once()

    def test_discovery_failure_is_not_silently_replaced_with_a_default(self):
        provider = AsyncMock()
        provider.list_models.side_effect = RuntimeError("network failure")
        with patch("product_council.desktop_session.create_provider", return_value=provider), \
                self.assertRaises(RuntimeError):
            asyncio.run(discover_models(self.choice()))
        provider.close.assert_awaited_once()

    def test_azure_discovery_requires_organization_deployment_not_catalog_model(self):
        with patch("product_council.desktop_session.create_provider") as create, \
                self.assertRaisesRegex(ModelDiscoveryError, "deployment name"):
            asyncio.run(discover_models(ModelChoice("azure_openai", "", "key", "https://example.com")))
        create.assert_not_called()

    def test_invalid_connection_response_never_passes_validation(self):
        for response in ('{"ok": false}', '{"message": "denied"}', "not JSON"):
            provider = AsyncMock()
            provider.complete.return_value = response
            with self.subTest(response=response), \
                    patch("product_council.desktop_session.create_provider", return_value=provider), \
                    self.assertRaises(ValueError):
                asyncio.run(validate_choice(self.choice()))
            provider.close.assert_awaited_once()

    def test_authentication_failure_is_not_swallowed(self):
        provider = AsyncMock()
        provider.complete.side_effect = RuntimeError("denied")
        with patch("product_council.desktop_session.create_provider", return_value=provider), \
                self.assertRaises(RuntimeError):
            asyncio.run(validate_choice(self.choice()))
        provider.close.assert_awaited_once()

    def test_session_uses_selected_model_for_every_seat_and_prompts_for_evidence(self):
        requests, research, events = [], [], []

        class Provider(MockProvider):
            async def complete(self, request):
                requests.append(request)
                return await super().complete(request)

            close = AsyncMock()

        async def clarify(questions):
            return {q: "Unknown" for q in questions}

        async def ask_research(agenda):
            research.append(agenda)
            return ("Customer interviews: renewal retention adoption are at risk. Internal review: engineering effort "
                    "and reuse; team vision roadmap; competitor competitive assessment; security misconfiguration "
                    "and secure score; support case volume. Source: customer-supplied pilot review, October 2026.")

        provider = Provider()
        with tempfile.TemporaryDirectory() as tmp, \
                patch("product_council.desktop_session.create_provider", return_value=provider), \
                patch("pathlib.Path.cwd", return_value=Path(tmp)):
            outcome = asyncio.run(run_session(
                self.choice(), QUESTION, _brief(), clarify=clarify, ask_research=ask_research,
                on_event=lambda kind, data: events.append((kind, data)),
            ))
            self.assertEqual(list(Path(tmp).iterdir()), [])
        self.assertEqual(outcome.status, "consensus")
        self.assertTrue(research)
        self.assertEqual(outcome.final_consensus.total, 8)
        self.assertTrue(all(r.model == self.choice().model for r in requests))
        self.assertEqual({r.metadata["role"] for r in requests},
                         {"customer_success", "engineering", "product", "marketing", "security", "support",
                          "forecast", "skeptic", "moderator"})
        self.assertTrue(all(e.source_type in ("user_clarification", "user_provided") for e in outcome.evidence))
        self.assertNotIn(self.choice().api_key, str(outcome.to_dict()))
        provider.close.assert_awaited_once()

    def test_provider_errors_cannot_leak_api_key_to_council(self):
        provider = AsyncMock()
        provider.complete.side_effect = RuntimeError("Authorization: Bearer test-key-not-a-real-secret")
        with self.assertRaises(RuntimeError) as caught:
            asyncio.run(DesktopProvider(provider).complete(LLMRequest("m", "s", "u")))
        self.assertNotIn(self.choice().api_key, str(caught.exception))
        self.assertIn("request failed", str(caught.exception))

    def test_safe_error_explains_authentication_model_and_quota_errors(self):
        for status, text in ((401, "Access denied"), (403, "Access denied"), (404, "not found"), (429, "quota")):
            exc = RuntimeError("secret response")
            exc.status_code = status
            self.assertIn(text, safe_error(exc))
            self.assertNotIn("secret response", safe_error(exc))


class WordReportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.outcome = _run(Council(_config(), retry_backoff=0))

    def test_no_file_or_directory_is_created_without_consent(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "proposal.docx"
            with self.assertRaisesRegex(ValueError, "approval"):
                save_word(self.outcome, path, approved=False)
            self.assertEqual(list(Path(tmp).iterdir()), [])

    def test_approved_document_contains_proposal_forecast_and_complete_audit(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "proposal.docx"
            save_word(self.outcome, path, approved=True)
            with ZipFile(path) as archive:
                self.assertIsNone(archive.testzip())
            document = Document(path)
            text = "\n".join(p.text for p in document.paragraphs)
            for value in (self.outcome.question, self.outcome.brief.hypothesis, self.outcome.brief.expected_result,
                          self.outcome.final_proposal.summary, "Forecast Agent", "Final votes"):
                self.assertIn(value, text)
            for metric in self.outcome.brief.success_metrics:
                self.assertIn(metric, text)
            for e in self.outcome.evidence:
                self.assertIn(e.excerpt, text)
            for claim in self.outcome.claims:
                self.assertIn(claim.text, text)
            for rec in self.outcome.iterations:
                for statement in rec.statements:
                    self.assertIn(statement.summary, text)
                for vote in rec.votes:
                    self.assertIn(vote.rationale, text)
            self.assertEqual(list(Path(tmp).iterdir()), [path])

    def test_failed_save_preserves_existing_file_and_removes_partial_output(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "proposal.docx"
            path.write_bytes(b"existing")
            with patch("docx.document.Document.save", side_effect=OSError("disk full")), self.assertRaises(OSError):
                save_word(self.outcome, path, approved=True)
            self.assertEqual(path.read_bytes(), b"existing")
            self.assertEqual(list(Path(tmp).iterdir()), [path])

    def test_no_consensus_remains_explicit_in_export(self):
        from docx import Document

        outcome = _run(Council(_config(), sources=[], retry_backoff=0))
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "no-consensus.docx"
            save_word(outcome, path, approved=True)
            text = "\n".join(p.text for p in Document(path).paragraphs)
            self.assertIn("NO CONSENSUS", text)
            self.assertIn(outcome.stop_reason, text)


class DesktopWindowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.outcome = _run(Council(_config(), retry_backoff=0))

    def setUp(self):
        import tkinter as tk
        from product_council.desktop import DesktopApp

        try:
            self.root = tk.Tk()
        except tk.TclError as exc:
            self.skipTest(f"Desktop display unavailable: {exc}")
        self.root.withdraw()
        self.store = ConnectionStore(MemoryVault())
        self.app = DesktopApp(self.root, store=self.store)
        self.root.update_idletasks()

    def tearDown(self):
        if hasattr(self, "root"):
            self.app.busy = False
            self.root.update_idletasks()
            self.app.close()

    def wait_for(self, predicate):
        deadline = time.monotonic() + 10
        while not predicate() and time.monotonic() < deadline:
            self.root.update()
            time.sleep(0.01)
        self.assertTrue(predicate(), "Desktop worker did not complete in time")

    def test_opt_in_save_and_future_launch_resume_without_key_entry(self):
        self.root.update()
        self.app.key_value.set("remembered-test-key")
        self.app.model.set("chosen-model")
        self.app.remember.set(True)
        with patch("product_council.desktop.validate_choice", new_callable=AsyncMock):
            self.app.connect()
            self.wait_for(lambda: hasattr(self.app, "idea_text"))
        saved = self.store.load()
        self.assertEqual(saved.api_key, "remembered-test-key")
        self.assertEqual(saved.model, "chosen-model")
        from product_council.desktop import DesktopApp

        self.app.close()
        import tkinter as tk
        self.root = tk.Tk()
        self.root.withdraw()
        with patch("product_council.desktop.validate_choice", new_callable=AsyncMock) as validate:
            self.app = DesktopApp(self.root, store=self.store)
            self.wait_for(lambda: hasattr(self.app, "idea_text"))
        self.assertEqual(validate.call_args.args[0], saved)
        self.assertEqual(self.app.choice, saved)

    def test_session_only_does_not_save_credentials(self):
        self.root.update()
        self.app.key_value.set("temporary-test-key")
        self.app.model.set("chosen-model")
        with patch("product_council.desktop.validate_choice", new_callable=AsyncMock):
            self.app.connect()
            self.wait_for(lambda: hasattr(self.app, "idea_text"))
        self.assertIsNone(self.store.load())

    def test_saved_key_failure_returns_to_editable_connection_form(self):
        self.root.update()
        saved = ModelChoice("openai", "chosen-model", "revoked-test-key")
        self.store.save(saved)
        with patch("product_council.desktop.validate_choice", new_callable=AsyncMock,
                   side_effect=RuntimeError("revoked")), \
                patch("product_council.desktop.messagebox.showerror") as error:
            self.app.resume_connection()
            self.wait_for(lambda: not self.app.busy)
        error.assert_called_once()
        self.assertEqual(str(self.app.key["state"]), "normal")
        self.assertIsNone(self.app.choice)

    def test_forget_only_selected_provider_and_cancel_does_not_remove(self):
        self.root.update()
        openai = ModelChoice("openai", "gpt-model", "openai-test-key")
        anthropic = ModelChoice("anthropic", "claude-model", "anthropic-test-key")
        self.store.save(openai)
        self.store.save(anthropic)
        self.app.fill_connection(openai)
        with patch("product_council.desktop.messagebox.askyesno", return_value=False):
            self.app.forget_connection()
        self.assertEqual(self.store.load("openai"), openai)
        with patch("product_council.desktop.messagebox.askyesno", return_value=True):
            self.app.forget_connection()
        self.assertIsNone(self.store.load("openai"))
        self.assertEqual(self.store.load("anthropic"), anthropic)
        self.assertEqual(self.app.key.get(), "")

    def test_unchecking_remember_removes_previously_saved_connection(self):
        self.root.update()
        saved = ModelChoice("openai", "gpt-model", "key")
        self.store.save(saved)
        self.app.fill_connection(saved)
        self.app.remember.set(False)
        with patch("product_council.desktop.validate_choice", new_callable=AsyncMock):
            self.app.connect()
            self.wait_for(lambda: hasattr(self.app, "idea_text"))
        self.assertIsNone(self.store.load())

    def test_save_failure_does_not_pretend_connection_was_remembered(self):
        self.root.update()
        self.app.key_value.set("key")
        self.app.model.set("gpt-model")
        self.app.remember.set(True)
        with patch("product_council.desktop.validate_choice", new_callable=AsyncMock), \
                patch.object(self.store, "save", side_effect=ConnectionStoreError("Storage unavailable")), \
                patch("product_council.desktop.messagebox.showerror") as error:
            self.app.connect()
            self.wait_for(lambda: not self.app.busy)
        error.assert_called_once()
        self.assertFalse(hasattr(self.app, "idea_text"))
        self.assertEqual(str(self.app.connect_button["state"]), "normal")

    def test_switch_provider_loads_its_separate_saved_connection(self):
        self.root.update()
        choice = ModelChoice("anthropic", "claude-model", "anthropic-test-key")
        self.store.save(choice)
        self.app.key_value.set("openai-test-key")
        self.app.provider.set("Anthropic")
        self.app.provider_changed()
        self.assertEqual(self.app.key.get(), choice.api_key)
        self.assertEqual(self.app.model.get(), choice.model)
        self.assertTrue(self.app.remember.get())

    def test_catalog_preselects_a_model_without_typing_and_allows_selection(self):
        self.app.key.insert(0, "key")
        with patch("product_council.desktop.discover_models", new_callable=AsyncMock,
                   return_value=["gpt-newer", "gpt-other"]) as discover:
            self.app.load_models()
            self.wait_for(lambda: not self.app.busy)
        discover.assert_awaited_once()
        self.assertEqual(self.app.model.get(), "gpt-newer")
        self.assertEqual(tuple(self.app.model["values"]), ("gpt-newer", "gpt-other"))
        self.assertEqual(str(self.app.model["state"]), "readonly")
        self.app.model.set("gpt-other")
        with patch("product_council.desktop.validate_choice", new_callable=AsyncMock) as validate:
            self.app.connect()
            self.wait_for(lambda: hasattr(self.app, "idea_text"))
        self.assertEqual(validate.call_args.args[0].model, "gpt-other")

    def test_changing_key_or_provider_invalidates_loaded_selection(self):
        self.app.model.configure(values=["gpt-example"])
        self.app.model.set("gpt-example")
        self.app.key_value.set("different-key")
        self.assertEqual(self.app.model.get(), "")
        self.assertFalse(self.app.model["values"])
        self.app.provider.set("Azure OpenAI")
        self.app.provider_changed()
        self.assertTrue(self.app.manual.get())
        self.assertEqual(str(self.app.load_button["state"]), "disabled")
        self.assertIn("administrator", self.app.status["text"])
        self.app.model.set("custom-deployment")
        self.app.endpoint_value.set("https://new-resource.openai.azure.com")
        self.assertEqual(self.app.model.get(), "")
        self.app.provider.set("Anthropic")
        self.app.provider_changed()
        self.assertFalse(self.app.manual.get())
        self.assertEqual(str(self.app.model["state"]), "readonly")
        self.assertEqual(str(self.app.load_button["state"]), "normal")

    def test_catalog_failure_preserves_credentials_and_allows_retry(self):
        self.app.key_value.set("key")
        with patch("product_council.desktop.discover_models", new_callable=AsyncMock,
                   side_effect=ModelDiscoveryError("No models available.")), \
                patch("product_council.desktop.messagebox.showerror") as error:
            self.app.load_models()
            self.wait_for(lambda: not self.app.busy)
        self.assertEqual(self.app.key.get(), "key")
        self.assertEqual(self.app.model.get(), "")
        self.assertEqual(str(self.app.load_button["state"]), "normal")
        error.assert_called_once()

    def test_wizard_rejects_empty_idea_and_brief(self):
        self.app.idea_page()
        with patch("product_council.desktop.messagebox.showerror") as error:
            self.app.accept_idea()
            error.assert_called_once()
        self.app.idea_text.insert("1.0", "A product idea")
        self.app.accept_idea()
        with patch("product_council.desktop.messagebox.showerror") as error:
            self.app.start_council()
            error.assert_called_once()
        self.assertFalse(self.app.busy)

    def test_preview_does_not_export_and_save_dialog_cancel_does_not_export(self):
        self.app.outcome = self.outcome
        with patch("product_council.desktop.save_word") as save:
            self.app.preview_page()
            save.assert_not_called()
            with patch("product_council.desktop.filedialog.asksaveasfilename", return_value=""):
                self.app.export()
            save.assert_not_called()

    def test_approval_saves_to_selected_path_only(self):
        self.app.outcome = _run(Council(_config(), retry_backoff=0))
        with tempfile.TemporaryDirectory() as tmp, \
                patch("product_council.desktop.filedialog.asksaveasfilename",
                      return_value=str(Path(tmp) / "approved.docx")), \
                patch("product_council.desktop.messagebox.showinfo") as notice:
            self.app.preview_page()
            self.assertEqual(list(Path(tmp).iterdir()), [])
            self.app.export()
            self.assertTrue((Path(tmp) / "approved.docx").is_file())
            notice.assert_called_once()

    def test_declining_preview_does_not_save(self):
        self.app.outcome = _run(Council(_config(), retry_backoff=0))
        with patch("product_council.desktop.messagebox.askyesno", return_value=True), \
                patch("product_council.desktop.save_word") as save:
            self.app.discard()
        self.assertIsNone(self.app.outcome)
        save.assert_not_called()

    def test_complete_desktop_flow_with_followups_and_approved_word_export(self):
        class Provider(MockProvider):
            async def list_models(self):
                return ["chosen-model"]

            async def complete(self, request):
                if not request.metadata:
                    return '{"ok": true}'
                return await super().complete(request)

        questions = []

        def answer(title, prompt, future):
            questions.append((title, prompt))
            if title == "Additional evidence needed":
                future.set_result("Source: customer pilot review. renewal retention adoption; engineering effort reuse; "
                                  "vision roadmap; competitor competitive; misconfiguration secure score; case volume.")
            else:
                future.set_result("Source: requester. We need to research the answer.")

        with tempfile.TemporaryDirectory() as tmp, \
                patch("product_council.desktop_session.create_provider", side_effect=lambda _: Provider()), \
                patch.object(self.app, "question_dialog", side_effect=answer), \
                patch("product_council.desktop.messagebox.showerror") as error:
            self.app.key.insert(0, "test-key")
            self.app.load_models()
            self.wait_for(lambda: not self.app.busy)
            self.assertEqual(self.app.model.get(), "chosen-model")
            self.app.connect()
            self.wait_for(lambda: hasattr(self.app, "idea_text"))
            self.app.idea_text.insert("1.0", QUESTION)
            self.app.accept_idea()
            brief = _brief()
            self.app.hypothesis.insert("1.0", brief.hypothesis)
            self.app.expected.insert("1.0", brief.expected_result)
            self.app.metrics.insert("1.0", "\n".join(brief.success_metrics))
            self.app.start_council()
            self.wait_for(lambda: self.app.outcome is not None)
            self.assertTrue(any(title == "Additional evidence needed" for title, _ in questions))
            self.assertEqual(self.app.outcome.status, "consensus")
            self.assertIsNone(self.app.choice)
            self.assertEqual(list(Path(tmp).iterdir()), [])
            with patch("product_council.desktop.filedialog.asksaveasfilename",
                       return_value=str(Path(tmp) / "approved.docx")), \
                    patch("product_council.desktop.messagebox.showinfo"):
                self.app.export()
            self.assertTrue((Path(tmp) / "approved.docx").is_file())
            error.assert_not_called()


class ProviderCatalogTests(unittest.TestCase):
    class Page:
        def __init__(self, models):
            self.models = models

        def __aiter__(self):
            async def items():
                for model in self.models:
                    yield model
            return items()

    def test_openai_filters_non_chat_models_and_sorts_and_deduplicates(self):
        from product_council.providers.openai_provider import OpenAIProvider

        provider = OpenAIProvider.__new__(OpenAIProvider)
        provider.name = "openai"
        models = [
            SimpleNamespace(id="gpt-old", created=10),
            SimpleNamespace(id="text-embedding-example", created=100),
            SimpleNamespace(id="gpt-new", created=20),
            SimpleNamespace(id="gpt-new", created=20),
            SimpleNamespace(id="gpt-audio", created=100),
            SimpleNamespace(id="gpt-realtime", created=100),
            SimpleNamespace(id="gpt-codex", created=100),
            SimpleNamespace(id="gpt-example-pro", created=100),
            SimpleNamespace(id="o3-example", created=15),
        ]
        client = AsyncMock()
        client.models.list.return_value = self.Page(models)
        provider._client = client
        self.assertEqual(asyncio.run(provider.list_models()), ["gpt-new", "o3-example", "gpt-old"])
        client.models.list.assert_awaited_once()
        provider.name = "azure_openai"
        with self.assertRaises(NotImplementedError):
            asyncio.run(provider.list_models())

    def test_anthropic_sorts_models_by_release_timestamp(self):
        from product_council.providers.anthropic_provider import AnthropicProvider

        provider = AnthropicProvider.__new__(AnthropicProvider)
        client = AsyncMock()
        client.models.list.return_value = self.Page([
            SimpleNamespace(id="claude-old", created_at=datetime(2025, 1, 1, tzinfo=timezone.utc)),
            SimpleNamespace(id="claude-new", created_at=datetime(2026, 1, 1, tzinfo=timezone.utc)),
        ])
        provider._client = client
        self.assertEqual(asyncio.run(provider.list_models()), ["claude-new", "claude-old"])


if __name__ == "__main__":
    unittest.main()
