"""Word export of the full council result. Called only after preview approval."""
from __future__ import annotations

import os
import tempfile
from pathlib import Path
from typing import Any

from .roles import ROLES
from .schemas import CouncilOutcome


def save_word(outcome: CouncilOutcome, path: Path, *, approved: bool) -> None:
    if not approved:
        raise ValueError("Explicit approval of the result preview is required before Word export.")
    from docx import Document

    document = Document()
    document.add_heading("Product Council proposal", 0)
    document.add_paragraph("Advisory decision for human review, not approval to release a product.")
    document.add_heading("Decision summary", 1)
    document.add_paragraph(f"Outcome: {outcome.status.replace('_', ' ').upper()}")
    document.add_paragraph(f"Reason: {outcome.stop_reason}")
    if outcome.final_consensus:
        c = outcome.final_consensus
        document.add_paragraph(f"Final vote: {c.agree}/{c.total} ({c.ratio:.0%}); threshold: {c.threshold:.0%}")
    document.add_heading("Productization idea", 1)
    document.add_paragraph(outcome.question)
    document.add_heading("Requester-defined hypothesis and success criteria", 1)
    document.add_paragraph(f"Hypothesis: {outcome.brief.hypothesis}")
    document.add_paragraph(f"Expected results: {outcome.brief.expected_result}")
    for metric in outcome.brief.success_metrics:
        document.add_paragraph(metric, style="List Bullet")
    document.add_paragraph("These are hypotheses and targets, not evidence of achieved outcomes.")

    def details(data: dict[str, Any], level: int = 2) -> None:
        for key, value in data.items():
            label = key.replace("_", " ").capitalize()
            if isinstance(value, dict):
                document.add_heading(label, min(level, 9))
                details(value, level + 1)
            elif isinstance(value, list):
                document.add_heading(label, min(level, 9))
                if not value:
                    document.add_paragraph("(none)")
                for index, item in enumerate(value, 1):
                    if isinstance(item, dict):
                        document.add_heading(f"{label} {index}", min(level + 1, 9))
                        details(item, level + 2)
                    else:
                        document.add_paragraph(str(item), style="List Bullet")
            else:
                document.add_paragraph(f"{label}: {value if value is not None else '(none)'}")

    raw = outcome.to_dict()
    document.add_heading("Final proposal", 1)
    if raw["final_proposal"]:
        details(raw["final_proposal"])
    else:
        document.add_paragraph("No proposal was produced.")
    document.add_heading("Forecast Agent: post-release outlook", 1)
    forecasts = [s for rec in outcome.iterations for s in rec.statements if s.role == "forecast"]
    for statement in forecasts:
        document.add_heading(f"Iteration {statement.iteration}", 2)
        document.add_paragraph(statement.summary)
    if not forecasts:
        document.add_paragraph("No forecast was produced. Review agent errors in the audit trail.")
    document.add_heading("Final votes", 1)
    if outcome.iterations:
        for vote in outcome.iterations[-1].votes:
            document.add_heading(ROLES[vote.role].title if vote.role in ROLES else vote.role, 2)
            details(vote.__dict__, 3)
    document.add_heading("Clarifications", 1)
    for question, answer in outcome.clarifications.items():
        document.add_paragraph(f"Question: {question}")
        document.add_paragraph(f"Answer: {answer}")
    document.add_heading("Full deliberation and research audit", 1)
    details({"iterations": raw["iterations"], "claims": raw["claims"]})
    document.add_heading("Evidence ledger (including source excerpts)", 1)
    details({"evidence": raw["evidence"]})

    # Write beside the destination, then replace it so a failed export cannot leave a partial report.
    fd, temporary = tempfile.mkstemp(suffix=".docx", dir=path.parent)
    os.close(fd)
    try:
        document.save(temporary)
        os.replace(temporary, path)
    finally:
        Path(temporary).unlink(missing_ok=True)
